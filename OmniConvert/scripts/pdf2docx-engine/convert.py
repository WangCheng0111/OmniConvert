# pdf2docx 引擎启动器(OmniConvert 内置引擎,个人学习自用)
# 用法: python.exe convert.py <输入.pdf> <输出.docx>
#
# 功能:
#   1. pdf2docx 基础转换(PDF -> 可编辑 DOCX);
#   2. 后处理一:清除"数字与中文之间"的多余半角空格
#      (PyMuPDF 提取层在数字与 CJK 跨字体片段间插入空格,原文档并无此空格);
#   3. 后处理二:恢复伪粗体
#      (宋体等无粗体字面的字体,Word 导出 PDF 时用描边伪粗体渲染;
#       PyMuPDF texttrace 中 type=1 的 span 即描边文本,按行聚合后
#       与 DOCX run 文本做规范化精确匹配,命中即置粗体)。
import os
import re
import sys
import zipfile
from collections import defaultdict

import fitz
from docx import Document
from pdf2docx import Converter

SPACE_RE = re.compile(r"\s+")
DIGIT_CJK_SPACE_RE = re.compile(r"(?<=[0-9]) (?=[\u4e00-\u9fff])")
THUMBNAIL_RELATIONSHIP_RE = re.compile(r"<Relationship\b[^>]*metadata/thumbnail[^>]*/>")


def normalize(text):
    return SPACE_RE.sub("", text)


def build_bold_regions(pdf_path):
    """按 (页, 行) 聚合描边(type=1)文本碎片,返回规范化后的粗体区域文本集合。"""
    lines = defaultdict(list)
    document = fitz.open(pdf_path)
    try:
        for page_index in range(len(document)):
            for span in document[page_index].get_texttrace():
                if span.get("type") != 1:
                    continue
                bbox = span.get("bbox")
                if not bbox:
                    continue
                text = "".join(chr(ch[0]) for ch in span["chars"])
                y_center = round((bbox[1] + bbox[3]) / 2, 0)
                lines[(page_index, y_center)].append((bbox[0], text))
    finally:
        document.close()

    regions = set()
    for fragments in lines.values():
        fragments.sort(key=lambda item: item[0])
        merged = normalize("".join(text for _, text in fragments))
        if len(merged) >= 2:
            regions.add(merged)
    return regions


def iter_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def post_process(pdf_path, docx_path):
    regions = build_bold_regions(pdf_path)
    document = Document(docx_path)
    fixed_spaces = 0
    bold_runs = 0
    for paragraph in iter_paragraphs(document):
        for run in paragraph.runs:
            original = run.text
            if original:
                cleaned = DIGIT_CJK_SPACE_RE.sub("", original)
                if cleaned != original:
                    run.text = cleaned
                    fixed_spaces += 1
            if regions and run.text and normalize(run.text) in regions:
                run.bold = True
                bold_runs += 1
    document.save(docx_path)
    return fixed_spaces, bold_runs


def strip_thumbnail(docx_path):
    """移除 python-docx 默认模板自带的 docProps/thumbnail.jpeg 及其关系。

    Windows 资源管理器在"中等"及以上图标尺寸会优先显示文档内嵌缩略图,
    模板自带的陈旧预览图会让转换产物显示"奇怪的 Word 图标";删除后与
    普通 Word 保存的文档一致(显示标准 Word 图标)。
    """
    with zipfile.ZipFile(docx_path, "r") as source:
        items = [(info, source.read(info.filename)) for info in source.infolist()]

    temp_path = docx_path + ".thumbtmp"
    with zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, data in items:
            if info.filename == "docProps/thumbnail.jpeg":
                continue
            if info.filename == "_rels/.rels":
                data = THUMBNAIL_RELATIONSHIP_RE.sub("", data.decode("utf-8")).encode("utf-8")
            target.writestr(info, data)
    os.replace(temp_path, docx_path)


def main():
    if len(sys.argv) < 3:
        sys.stderr.write("usage: convert.py <input.pdf> <output.docx>\n")
        return 2

    input_path, output_path = sys.argv[1], sys.argv[2]

    converter = Converter(input_path)
    try:
        converter.convert(output_path)
    finally:
        converter.close()

    fixed_spaces, bold_runs = post_process(input_path, output_path)
    strip_thumbnail(output_path)
    print("post-process: spaces={0}, bold={1}".format(fixed_spaces, bold_runs))
    return 0


if __name__ == "__main__":
    sys.exit(main())
